import io
import re
import pymupdf
from docx import Document as DocxFile
from langchain_core.documents import Document


# every loader returns a list of Documents with metadata:
#   paper -> file name,  page -> number used for ordering,  loc -> label shown in citations


def fix_hyphens(text):
    # join words split at line ends in PDFs: "retrie-\nval" -> "retrieval"
    return re.sub(r"(\w)-\n(\w)", r"\1\2", text)


def load_pymupdf(data, paper):
    pages = []
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        for i, page in enumerate(pdf, start=1):
            pages.append(Document(page_content=fix_hyphens(page.get_text()),
                                  metadata={"paper": paper, "page": i, "loc": f"p. {i}"}))
    return pages


def load_docx(data, paper):
    # Word files have no fixed pages, so each heading starts a new block
    # and citations point to the heading (or paragraph number if there is no heading)
    doc = DocxFile(io.BytesIO(data))
    blocks, lines, loc = [], [], "para 1"

    for i, para in enumerate(doc.paragraphs, start=1):
        text = para.text.strip()
        if not text:
            continue
        if para.style.name.lower().startswith(("heading", "title")):
            if lines:
                blocks.append((loc, lines))
            lines, loc = [], text[:60]
        lines.append(text)
    if lines:
        blocks.append((loc, lines))

    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        blocks.append(("tables", rows))

    return [Document(page_content="\n".join(lines),
                     metadata={"paper": paper, "page": i, "loc": loc})
            for i, (loc, lines) in enumerate(blocks, start=1)]


def load_file(data, file_name):
    # loader is picked from the file type: PDF -> PyMuPDF, DOCX -> python-docx
    paper = file_name.rsplit(".", 1)[0]
    if file_name.lower().endswith(".docx"):
        pages = load_docx(data, paper)
    else:
        pages = load_pymupdf(data, paper)

    # pages with (almost) no text are usually scanned images -> no OCR yet
    empty = [p.metadata["loc"] for p in pages if len(p.page_content.strip()) < 20]
    pages = [p for p in pages if len(p.page_content.strip()) >= 20]
    return paper, pages, empty
